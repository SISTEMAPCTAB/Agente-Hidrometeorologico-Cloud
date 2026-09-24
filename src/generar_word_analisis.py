import os
import hashlib
import json
import re
import shutil
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE = Path("data/analysis")
SOURCE = Path(os.environ.get("ANALYSIS_SOURCE", str(BASE / "latest.md")))
PUBLIC = Path("data/fuentes_publicas/latest.json")
HIST = Path("data/fuentes_publicas/insivumeh_historico.json")
LEVELS = Path("data/niveles/Ultimo_Corte/ultimo_resumen.json")
REPORT_DIR = BASE / "Reportes"
ARCHIVE_DIR = REPORT_DIR / "Archivo"
LATEST = REPORT_DIR / "Analisis_Hidrometeorologico_Ultimo.docx"
WORD_STATUS = REPORT_DIR / "Estado_Word.json"
BLUE = RGBColor(31, 78, 121)

def clean_inline(text):
    text=text.strip(); text=re.sub(r"\*\*(.*?)\*\*",r"\1",text); text=re.sub(r"__(.*?)__",r"\1",text); return re.sub(r"`([^`]+)`",r"\1",text)
def extract_cut(text):
    m=re.search(r"(?im)^\s*(?:Corte|CORTE)\s*:\s*(\d{4}-\d{2}-\d{2})(?:\s+|T)(\d{1,2}:\d{2})?",text)
    return datetime.fromisoformat(f"{m.group(1)}T{m.group(2) or '11:00'}") if m else datetime.now()
def load_json(p,default):
    try:return json.loads(p.read_text(encoding='utf-8'))
    except:return default
def n(v):
    try:return float(v)
    except:return None
def add_heading(doc,text,level):
    p=doc.add_paragraph(); r=p.add_run(clean_inline(text)); r.font.name='Arial'; r.font.bold=True; r.font.color.rgb=BLUE; r.font.size=Pt({1:14,2:12,3:10.5}.get(level,10)); return p
def add_table(doc,lines):
    rows=[[clean_inline(x) for x in line.strip().strip('|').split('|')] for line in lines]
    if len(rows)>=2 and all(re.fullmatch(r"\s*:?-{3,}:?\s*",c or '') for c in rows[1]): rows.pop(1)
    if not rows:return
    cols=max(len(r) for r in rows); table=doc.add_table(rows=0,cols=cols); table.style='Table Grid'
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for ci in range(cols):
            cells[ci].text=row[ci] if ci<len(row) else ''
            for p in cells[ci].paragraphs:
                for run in p.runs: run.font.name='Arial'; run.font.size=Pt(8); run.font.bold=ri==0
def render_markdown(doc,text):
    lines=text.splitlines(); i=0; buf=[]
    def flush():
        nonlocal buf
        joined=' '.join(clean_inline(x) for x in buf if x.strip()).strip()
        if joined:
            p=doc.add_paragraph(joined)
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(9.5)
        buf=[]
    while i<len(lines):
        line=lines[i].rstrip()
        if not line.strip():flush();i+=1;continue
        if line.startswith('### '):flush();add_heading(doc,line[4:],3);i+=1;continue
        if line.startswith('## '):flush();add_heading(doc,line[3:],2);i+=1;continue
        if line.startswith('# '):flush();add_heading(doc,line[2:],1);i+=1;continue
        if line.lstrip().startswith('|') and '|' in line.strip('|'):
            flush(); block=[]
            while i<len(lines) and lines[i].lstrip().startswith('|'):block.append(lines[i]);i+=1
            add_table(doc,block);continue
        if re.match(r'^\s*[-*]\s+',line):
            flush();p=doc.add_paragraph(style='List Bullet');p.add_run(clean_inline(re.sub(r'^\s*[-*]\s+','',line)))
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(9.5)
            i+=1;continue
        buf.append(line);i+=1
    flush()
def insivumeh_block():
    d=load_json(PUBLIC,{}); h=load_json(HIST,{}); levels=load_json(LEVELS,[])
    actual=d.get('insivumeh_bhvgm_actual') or {}; sts=actual.get('estaciones') or []
    def st(name):return next((x for x in sts if x.get('estacion')==name),None)
    def fmt(x,k,u):
        v=n(x.get(k)) if x else None; return 's/d' if v is None else f'{v:.2f} {u}'
    por=st('El Porvenir'); tig=st('El Tigre'); ant=d.get('insivumeh_el_tigre_antecedente')
    def lev(name):
        q=name.lower();return next((x for x in levels if q in f"{x.get('estacion','')} {x.get('rio','')}".lower()),None)
    chain=[]
    for label,key in [('Boca del Cerro','boca del cerro'),('Balancán','balanc'),('Emiliano Zapata','emiliano'),('Jonuta','jonuta'),('Centla/Tres Brazos','tres brazos')]:
        x=lev(key); lv=n(x.get('ultimo_nivel')) if x else None; de=n(x.get('delta_24h')) if x else None
        chain.append(f"{label}: "+('s/d' if lv is None else f'{lv:.2f} m')+((' · Δ24 h '+f'{de:+.2f} m') if de is not None else ''))
    lines=['\n## Cuenca alta del Usumacinta · señal observada INSIVUMEH','La interpretación mantiene separados los cortes y no atribuye causalidad ni usa el desfase horario como tiempo de tránsito. Los máximos BHVGM son referencias estadísticas, no umbrales de inundación.']
    lines.append(f"- El Porvenir–La Pasión: nivel {fmt(por,'nivel_instantaneo_m','m')} · caudal {fmt(por,'caudal_instantaneo_m3s','m³/s')}.")
    if tig: lines.append(f"- El Tigre–Usumacinta: nivel {fmt(tig,'nivel_instantaneo_m','m')} · caudal {fmt(tig,'caudal_instantaneo_m3s','m³/s')}.")
    elif ant:
        a=ant.get('estacion') or {}; lines.append(f"- El Tigre–Usumacinta: no publicado en el BHVGM actual. Último antecedente oficial: {ant.get('fecha','s/d')} {ant.get('hora','s/d')} h, nivel {fmt(a,'nivel_instantaneo_m','m')} · caudal {fmt(a,'caudal_instantaneo_m3s','m³/s')}. Antecedente informativo: no sustituye el dato actual, no genera tendencia y no gobierna semáforo.")
    else: lines.append('- El Tigre–Usumacinta: sin dato publicado utilizable para el corte; no se hereda ni se fabrica valor.')
    lines.append('- Cadena de seguimiento en Tabasco: '+' → '.join(chain)+'.')
    ev=d.get('insivumeh_evolucion_real') or []; et=next((x for x in ev if x.get('estacion')=='El Tigre'),None)
    lines.append('- Evolución El Tigre: '+(f"Δ nivel {n(et.get('delta_nivel_m')):+.3f} m en {n(et.get('horas_transcurridas')):.1f} h; tendencia {et.get('tendencia_calculada','s/d')}." if et and n(et.get('delta_nivel_m')) is not None and n(et.get('horas_transcurridas')) is not None else 'sin tendencia comparable; se requieren al menos dos cortes válidos que publiquen la estación.'))
    lines.append('- Lectura operativa: Chixoy/Salinas + La Pasión/El Porvenir + Usumacinta/El Tigre → Lacantún → Boca del Cerro → Balancán → Emiliano Zapata → Jonuta → Centla. Los tiempos históricos de tránsito solo se usan como contexto cuando existen observaciones compatibles; nunca de forma determinista.')
    return '\n'.join(lines)+'\n'
def main():
    REPORT_DIR.mkdir(parents=True,exist_ok=True);ARCHIVE_DIR.mkdir(parents=True,exist_ok=True)
    if not SOURCE.exists() or SOURCE.stat().st_size<100:
        status={'report':'Análisis hidrometeorológico diario','verified':False,'state':'waiting_for_latest_md'};WORD_STATUS.write_text(json.dumps(status,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');return 0
    text=SOURCE.read_text(encoding='utf-8-sig'); cut=extract_cut(text); archive_name=f"Analisis_Hidrometeorologico_{cut.strftime('%Y-%m-%d')}.docx"
    if SOURCE == BASE / 'latest.md' and 'Cuenca alta del Usumacinta · señal observada INSIVUMEH' not in text:text=text.rstrip()+insivumeh_block()
    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.55);sec.left_margin=Inches(.6);sec.right_margin=Inches(.6);doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(9.5)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Actualización del análisis operativo hidrometeorológico');r.font.name='Arial';r.font.size=Pt(16);r.font.bold=True;r.font.color.rgb=BLUE
    p2=doc.add_paragraph();p2.alignment=WD_ALIGN_PARAGRAPH.CENTER;r2=p2.add_run('Cuencas Grijalva–Usumacinta');r2.font.name='Arial';r2.font.size=Pt(10);r2.font.bold=True
    render_markdown(doc,text);doc.save(LATEST)
    with zipfile.ZipFile(LATEST,'r') as zf:
        bad=zf.testzip()
        if bad:raise RuntimeError(f'DOCX corrupto; entrada dañada: {bad}')
    Document(str(LATEST))
    if LATEST.stat().st_size<5000:raise RuntimeError('DOCX generado demasiado pequeño.')
    archived=ARCHIVE_DIR/archive_name;shutil.copy2(LATEST,archived);digest=hashlib.sha256(LATEST.read_bytes()).hexdigest()
    status={'report':'Análisis hidrometeorológico diario','cut':cut.isoformat(),'verified':True,'source_path':str(SOURCE).replace('\\','/'),'latest_path':str(LATEST).replace('\\','/'),'archive_path':str(archived).replace('\\','/'),'bytes':LATEST.stat().st_size,'sha256':digest,'insivumeh_usumacinta_integrado':True}
    WORD_STATUS.write_text(json.dumps(status,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(status,ensure_ascii=False,indent=2));return 0
if __name__=='__main__':sys.exit(main())
