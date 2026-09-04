import hashlib
import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE = Path("data/niveles")
SUMMARY = BASE / "Ultimo_Corte" / "ultimo_resumen.json"
STATUS_SOURCE = BASE / "Estado_Niveles_Cloud.json"
GRAPH_DIR = BASE / "Graficas"
REPORT_DIR = BASE / "Reportes"
ARCHIVE_DIR = REPORT_DIR / "Archivo"
LATEST = REPORT_DIR / "Reporte_Niveles_Hidrometricos_Ultimo.docx"
WORD_STATUS = REPORT_DIR / "Estado_Word.json"

BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(45, 45, 45)


def load_json(path, default):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def parse_dt(value):
    if not value:
        return None
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def fmt(v, signed=False):
    if v is None:
        return "s/d"
    try:
        n = float(v)
    except Exception:
        return str(v)
    return f"{n:+.2f}" if signed else f"{n:.2f}"


def set_cell_font(cell, size=7.5, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(size)
            r.font.bold = bold


def style_table(table, header_size=7.3, body_size=7.0):
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_font(cell, header_size if i == 0 else body_size, i == 0)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.name = "Arial"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = BLUE
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = DARK


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12 if level == 1 else 10.5)
    r.font.bold = True
    r.font.color.rgb = BLUE
    return p


def add_dashboard(doc):
    dashboard = GRAPH_DIR / "tablero_operativo.png"
    if dashboard.exists() and dashboard.stat().st_size > 0:
        add_heading(doc, "Tablero gráfico operativo", 1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(dashboard), width=Inches(9.2))


def add_system_graphs(doc):
    system_dir = GRAPH_DIR / "Sistemas"
    if not system_dir.exists():
        return
    imgs = sorted(system_dir.glob("*.png"))
    if not imgs:
        return
    add_heading(doc, "Gráficas por sistema hidrológico", 1)
    for img in imgs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(img.stem.replace("_", " "))
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run().add_picture(str(img), width=Inches(9.0))


def main():
    summary = load_json(SUMMARY, [])
    source_status = load_json(STATUS_SOURCE, {})
    if not summary:
        raise RuntimeError("No existe resumen de niveles para generar Word.")

    cut = parse_dt(summary[0].get("fecha_hora")) or datetime.now()
    cut_label = cut.strftime("%d/%m/%Y %H:%M h")
    archive_name = f"Reporte_Niveles_Hidrometricos_{cut.strftime('%Y-%m-%d_%H%M')}.docx"

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.45)
    sec.right_margin = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    add_title(doc, "Vigilancia de Niveles Hidrométricos", f"Corte: {cut_label} | Ciclo operativo 06:00–06:00")

    p = doc.add_paragraph()
    p.add_run("Estado del agente: ").bold = True
    p.add_run("COMPLETO" if source_status.get("complete") else "INCOMPLETO")
    p.add_run(" | Regla de continuidad: ").bold = True
    p.add_run("si falta lectura nueva, se conserva el último valor válido y se marca como heredado; ausencia nunca equivale a cero.")

    add_heading(doc, "Resumen por estación", 1)
    headers = ["Sistema", "Estación", "Actual m", "Base 06 m", "Δ vs 06 m", "Δ anterior m", "Tendencia", "Dato", "Última observación real"]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in summary:
        cells = table.add_row().cells
        vals = [
            row.get("rio", ""), row.get("estacion", ""), fmt(row.get("ultimo_nivel")),
            fmt(row.get("nivel_base_06")), fmt(row.get("delta_vs_06"), True),
            fmt(row.get("delta_reporte"), True), row.get("tendencia", ""),
            row.get("estado_dato", ""), row.get("ultima_actualizacion_real", "")
        ]
        for i, v in enumerate(vals):
            cells[i].text = str(v)
    style_table(table)

    add_heading(doc, "Referencia NAMO/NAME y ventanas históricas", 1)
    headers2 = ["Estación", "Δ24 h", "Δ48 h", "Δ72 h", "Dist. NAMO m", "Dist. NAME m", "Base 06", "Estado base"]
    table2 = doc.add_table(rows=1, cols=len(headers2))
    for i, h in enumerate(headers2):
        table2.rows[0].cells[i].text = h
    for row in summary:
        cells = table2.add_row().cells
        vals = [
            row.get("estacion", ""), fmt(row.get("delta_24h"), True), fmt(row.get("delta_48h"), True),
            fmt(row.get("delta_72h"), True), fmt(row.get("distancia_namo"), True),
            fmt(row.get("distancia_name"), True), row.get("fecha_hora_base_ciclo", ""), row.get("estado_base_06", "")
        ]
        for i, v in enumerate(vals):
            cells[i].text = str(v)
    style_table(table2)

    add_dashboard(doc)
    add_system_graphs(doc)

    add_heading(doc, "Criterio operativo", 1)
    doc.add_paragraph(
        "La variación Δ vs 06:00 mide el cambio acumulado dentro del ciclo operativo vigente. "
        "La variación Δ anterior mide el cambio inmediato respecto al corte previo. Los valores heredados "
        "se muestran para continuidad operativa y no deben interpretarse como observaciones nuevas."
    )

    doc.save(LATEST)
    # Verificación estructural: ZIP Office válido y reapertura con python-docx.
    with zipfile.ZipFile(LATEST, "r") as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"DOCX corrupto; entrada dañada: {bad}")
    Document(str(LATEST))
    if LATEST.stat().st_size < 5000:
        raise RuntimeError("DOCX generado demasiado pequeño.")

    archived = ARCHIVE_DIR / archive_name
    shutil.copy2(LATEST, archived)
    digest = hashlib.sha256(LATEST.read_bytes()).hexdigest()

    status = {
        "report": "Vigilancia Niveles Hidrométricos",
        "cut": cut.isoformat(),
        "verified": True,
        "latest_path": str(LATEST).replace("\\", "/"),
        "archive_path": str(archived).replace("\\", "/"),
        "bytes": LATEST.stat().st_size,
        "sha256": digest,
        "graphs_embedded": (GRAPH_DIR / "tablero_operativo.png").exists(),
    }
    WORD_STATUS.write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
