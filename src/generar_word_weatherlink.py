import csv
import hashlib
import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE = Path("data/weatherlink")
CSV_PATH = BASE / "Tabla_WeatherLink_Cloud.csv"
STATUS_PATH = BASE / "Estado_WeatherLink_Cloud.json"
CONAGUA_BASE = Path("data/niveles/Lluvia_CONAGUA")
CONAGUA_STATUS = CONAGUA_BASE / "Estado_Lluvia_CONAGUA.json"
CONAGUA_LATEST = CONAGUA_BASE / "ultimo_corte.json"
REPORT_DIR = BASE / "Reportes"
ARCHIVE_DIR = REPORT_DIR / "Archivo"
LATEST = REPORT_DIR / "Reporte_WeatherLink_Ultimo.docx"
WORD_STATUS = REPORT_DIR / "Estado_Word.json"

BLUE = RGBColor(31, 78, 121)


def load_json(path, default):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def read_rows():
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def num(value):
    try:
        return float(str(value).strip())
    except Exception:
        return None


def set_font(cell, size=7.1, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(size)
            r.font.bold = bold


def style_table(table, header=7.2, body=6.8):
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_font(cell, header if i == 0 else body, i == 0)


def heading(doc, text, size=12):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = BLUE
    return p


def main():
    if not CSV_PATH.exists():
        raise RuntimeError("No existe Tabla_WeatherLink_Cloud.csv")
    status = load_json(STATUS_PATH, {})
    rows = read_rows()
    if not rows:
        raise RuntimeError("WeatherLink no tiene filas para reportar.")

    conagua_status = load_json(CONAGUA_STATUS, {})
    conagua_rows = load_json(CONAGUA_LATEST, [])

    cut = datetime.fromisoformat(status.get("run_local", datetime.now().astimezone().isoformat()).replace("Z", "+00:00"))
    cut_label = cut.strftime("%d/%m/%Y %H:%M h")
    archive_name = f"Reporte_WeatherLink_{cut.strftime('%Y-%m-%d_%H%M')}.docx"

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Inches(0.4)
    sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.4)
    sec.right_margin = Inches(0.4)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Vigilancia de Lluvia Observada")
    r.font.name = "Arial"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = BLUE
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"WeatherLink + complemento CONAGUA | Corte: {cut_label}")
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)

    p = doc.add_paragraph()
    p.add_run("WeatherLink cloud: ").bold = True
    p.add_run("COMPLETO" if status.get("complete") else "INCOMPLETO")
    p.add_run(f" | Estaciones OK: {status.get('stations_ok', 's/d')}/{status.get('stations_expected', 's/d')}")
    maturity = status.get("window_maturity", {})
    p.add_run(" | Madurez ventanas: ").bold = True
    p.add_run(", ".join(f"{k}={'sí' if v else 'no'}" for k, v in maturity.items()))

    heading(doc, "Máximos WeatherLink", 11.5)
    for col, label in [("1 h mm", "1 h"), ("24 h mm", "24 h")]:
        valid = [(num(r.get(col)), r) for r in rows]
        valid = [(v, r) for v, r in valid if v is not None]
        if valid:
            v, rr = max(valid, key=lambda x: x[0])
            p = doc.add_paragraph()
            p.add_run(f"Máximo {label}: ").bold = True
            p.add_run(f"{v:.1f} mm en {rr.get('Estacion')} ({rr.get('Corte fuente')}); calidad: {rr.get('Calidad')}.")

    delayed = [r for r in rows if (r.get("Estado") or "").lower() != "ok"]
    p = doc.add_paragraph()
    p.add_run("Estaciones WeatherLink con dato rezagado o no OK: ").bold = True
    p.add_run(str(len(delayed)))

    heading(doc, "WeatherLink — 32 estaciones", 11.5)
    headers = ["Estación", "Ubicación", "1 h", "6 h", "24 h", "48 h", "72 h", "Mes", "Rango 24 h", "Estado", "Corte fuente", "Calidad"]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for rr in rows:
        vals = [
            rr.get("Estacion", ""), rr.get("Ubicacion", ""), rr.get("1 h mm", ""), rr.get("6 h mm", ""),
            rr.get("24 h mm", ""), rr.get("48 h mm", ""), rr.get("72 h mm", ""), rr.get("Mes mm", ""),
            rr.get("Rango SMN 24h", ""), rr.get("Estado", ""), rr.get("Corte fuente", ""), rr.get("Calidad", "")
        ]
        cells = table.add_row().cells
        for i, v in enumerate(vals):
            cells[i].text = str(v)
    style_table(table)

    heading(doc, "Lluvia observada CONAGUA — complemento", 11.5)
    p = doc.add_paragraph()
    p.add_run("Regla de captura: ").bold = True
    p.add_run("celda en blanco = sin dato; nunca se convierte a 0 mm. Solo se muestra 0 mm cuando CONAGUA publica explícitamente cero.")

    if not conagua_rows:
        doc.add_paragraph("Sin corte CONAGUA disponible para complementar este reporte.")
    else:
        mode = conagua_status.get("modo", "s/d")
        p = doc.add_paragraph()
        p.add_run("Estado de identificación CONAGUA: ").bold = True
        p.add_run(mode)
        if mode == "columnas_no_hidrometricas_en_revision":
            doc.add_paragraph("Las columnas detectadas permanecen en validación y no se presentan como estaciones oficiales hasta confirmar sus encabezados.")

        headers_c = ["Hora", "Estación/columna", "Lluvia mm", "Estado", "Fuente"]
        table_c = doc.add_table(rows=1, cols=len(headers_c))
        for i, h in enumerate(headers_c):
            table_c.rows[0].cells[i].text = h
        for rr in conagua_rows:
            lluvia = rr.get("lluvia_mm")
            lluvia_txt = "s/d" if lluvia is None else f"{float(lluvia):.1f}"
            vals = [rr.get("fecha_hora", ""), rr.get("estacion", ""), lluvia_txt, rr.get("estado_dato", ""), rr.get("fuente", "CONAGUA")]
            cells = table_c.add_row().cells
            for i, v in enumerate(vals):
                cells[i].text = str(v)
        style_table(table_c, header=7.5, body=7.2)

        valid_c = [(num(r.get("lluvia_mm")), r) for r in conagua_rows]
        valid_c = [(v, r) for v, r in valid_c if v is not None]
        if valid_c:
            v, rr = max(valid_c, key=lambda x: x[0])
            p = doc.add_paragraph()
            p.add_run("Máximo CONAGUA del corte: ").bold = True
            p.add_run(f"{v:.1f} mm en {rr.get('estacion', 's/d')} ({rr.get('fecha_hora', 's/d')}).")

    heading(doc, "Criterio de interpretación", 11.5)
    doc.add_paragraph(
        "WeatherLink y CONAGUA se mantienen como redes independientes. No se promedian ni se corrige una con la otra. "
        "Los acumulados WeatherLink cuya ventana cloud aún no ha madurado o cuya calidad indique 'heredado' se consideran transitorios. "
        "En CONAGUA, una celda vacía significa sin dato y no representa ausencia de lluvia."
    )

    doc.save(LATEST)
    with zipfile.ZipFile(LATEST, "r") as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"DOCX corrupto; entrada dañada: {bad}")
    Document(str(LATEST))
    if LATEST.stat().st_size < 5000:
        raise RuntimeError("DOCX generado demasiado pequeño.")

    archived = ARCHIVE_DIR / archive_name
    shutil.copy2(LATEST, archived)
    digest = hashlib.sha256(LATEST.read_bytes()).hexdigest()
    word_status = {
        "report": "Vigilancia WeatherLink + CONAGUA",
        "cut": cut.isoformat(),
        "verified": True,
        "latest_path": str(LATEST).replace("\\", "/"),
        "archive_path": str(archived).replace("\\", "/"),
        "bytes": LATEST.stat().st_size,
        "sha256": digest,
        "conagua_included": bool(conagua_rows),
        "conagua_mode": conagua_status.get("modo"),
    }
    WORD_STATUS.write_text(json.dumps(word_status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(word_status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
