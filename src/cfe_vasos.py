"""Colector CFE · Funcionamiento de vasos del Sistema Hidroeléctrico Grijalva.

Lee de Gmail el correo más reciente de Hidrometría Grijalva (CFE), asunto "FUNCIONAMIENTO DE VASOS",
con adjunto FINddmmaa.xlsx, y publica sólo los valores horarios que usa el modelo:
Peñitas (nivel, aportado, turbinado, vertido, extracción), canal Juan de Grijalva, Malpaso y niveles de
Angostura y Chicoasén. El Excel no se guarda ni se publica.

Salidas: data/cfe_vasos/latest.json y data/cfe_vasos/historico/AAAA-MM-DD.json (última versión de cada día).
Estado: data/cfe_vasos/estado.json (último mensaje procesado; evita volver a descargar el mismo adjunto).
Nunca falla el flujo de Gmail: cualquier error se registra en estado.json y el script sale con 0.
"""
import email
import imaplib
import io
import json
import os
import re
import time
from datetime import datetime, timedelta, timezone
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from pathlib import Path

OUT = Path('data/cfe_vasos')
STATE = OUT / 'estado.json'
TZ = timezone(timedelta(hours=-6))
NAMO = {'Peñitas': 87.40, 'Malpaso': 182.50, 'Chicoasén': 392.50}
NAME_RE = re.compile(r'^FIN(\d{2})(\d{2})(\d{2})\.xlsx?$', re.I)


def now_local():
    return datetime.now(TZ)


def num(v):
    try:
        x = float(v)
        return x if x == x else None
    except (TypeError, ValueError):
        return None


def r2(v, d=2):
    return None if v is None else round(v, d)


def hour_of(v):
    if hasattr(v, 'hour'):
        return v.hour if v.hour else None
    m = re.search(r'(\d{1,2})\s*:\s*00', str(v or ''))
    return int(m.group(1)) if m else None


def parse_xlsx(blob):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    ws = wb['DatosJG']
    filas = []
    for row in ws.iter_rows(min_row=8, max_row=40, max_col=15, values_only=True):
        h = hour_of(row[0])
        if h is None and str(row[0] or '').strip().startswith('24'):
            h = 24
        if h is None:
            continue
        f = {'hora': h,
             'malpaso': {'nivel': num(row[1]), 'turbinado': num(row[2]), 'vertido': num(row[3]), 'extraccion': num(row[4])},
             'juan_grijalva': {'nivel': num(row[5]), 'aportado': num(row[6]), 'canal': num(row[7]), 'tuneles': num(row[8]), 'total': num(row[9])},
             'penitas': {'nivel': num(row[10]), 'aportado': num(row[11]), 'turbinado': num(row[12]), 'vertido': num(row[13]), 'extraccion': num(row[14])}}
        if f['penitas']['nivel'] is None and f['malpaso']['nivel'] is None:
            continue
        filas.append(f)
    if not filas:
        raise ValueError('DatosJG sin horas con dato')
    # Aportación de Malpaso y niveles de Angostura/Chicoasén (hojas Rep_Peñ y ElevSHG).
    apm, niv = {}, {}
    try:
        for row in wb['Rep_Peñ'].iter_rows(min_row=8, max_row=40, max_col=3, values_only=True):
            h = hour_of(row[0])
            if h and num(row[2]) is not None:
                apm[h] = num(row[2])
    except KeyError:
        pass
    try:
        for row in wb['ElevSHG'].iter_rows(min_row=5, max_row=30, max_col=6, values_only=True):
            h = int(row[0]) if isinstance(row[0], (int, float)) else None
            if h and num(row[1]) is not None:
                niv[h] = {'angostura': num(row[1]), 'chicoasen': num(row[2])}
    except KeyError:
        pass
    novedades = None
    try:
        for row in wb['DATOS'].iter_rows(min_row=1, max_row=30, max_col=10, values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip().startswith('Presa Malpaso') and len(v) > 80:
                    novedades = re.sub(r'\s+', ' ', v).strip()
    except KeyError:
        pass
    wb.close()
    return filas, apm, niv, novedades


def build(blob, fname, msg_dt):
    m = NAME_RE.match(fname)
    fecha = f'20{m.group(3)}-{m.group(2)}-{m.group(1)}' if m else msg_dt.astimezone(TZ).date().isoformat()
    filas, apm, niv, novedades = parse_xlsx(blob)
    pe = [f for f in filas if f['penitas']['extraccion'] is not None]
    ult = pe[-1] if pe else filas[-1]
    h = ult['hora']

    def prom(key, sub, n=None):
        xs = [f[key][sub] for f in filas if f[key][sub] is not None and (n is None or f['hora'] > h - n)]
        return r2(sum(xs) / len(xs), 1) if xs else None

    p, jg, mp = ult['penitas'], ult['juan_grijalva'], ult['malpaso']
    horario = [{'hora': f['hora'],
                'penitas_nivel': r2(f['penitas']['nivel']), 'penitas_aportado': r2(f['penitas']['aportado'], 1),
                'penitas_turbinado': r2(f['penitas']['turbinado'], 1), 'penitas_vertido': r2(f['penitas']['vertido'], 1),
                'penitas_extraccion': r2(f['penitas']['extraccion'], 1),
                'jg_nivel': r2(f['juan_grijalva']['nivel']), 'jg_aportado': r2(f['juan_grijalva']['aportado'], 1),
                'jg_salida': r2(f['juan_grijalva']['total'], 1),
                'malpaso_nivel': r2(f['malpaso']['nivel']), 'malpaso_extraccion': r2(f['malpaso']['extraccion'], 1),
                'malpaso_aportado': r2(apm.get(f['hora']), 1),
                'angostura_nivel': r2((niv.get(f['hora']) or {}).get('angostura')),
                'chicoasen_nivel': r2((niv.get(f['hora']) or {}).get('chicoasen'))} for f in filas]
    hora_iso = (datetime.fromisoformat(fecha).replace(tzinfo=TZ) + timedelta(hours=h)).isoformat()
    return {
        'producto': 'Funcionamiento de vasos · Sistema Hidroeléctrico Grijalva',
        'fuente': 'CFE · Subgerencia de Producción Hidroeléctrica Grijalva, Departamento Regional de Hidrometría',
        'fecha': fecha, 'hora_ultimo_dato': h, 'fecha_hora_ultimo_dato': hora_iso,
        'mensaje_local': msg_dt.astimezone(TZ).isoformat(timespec='minutes'),
        'actualizado_local': now_local().isoformat(timespec='minutes'),
        'penitas': {'nivel_msnm': r2(p['nivel']), 'namo_msnm': NAMO['Peñitas'],
                    'dif_namo_m': r2(NAMO['Peñitas'] - p['nivel']) if p['nivel'] is not None else None,
                    'aportado_m3s': r2(p['aportado'], 1), 'turbinado_m3s': r2(p['turbinado'], 1),
                    'vertido_m3s': r2(p['vertido'] or 0, 1), 'extraccion_m3s': r2(p['extraccion'], 1),
                    'extraccion_prom_6h_m3s': prom('penitas', 'extraccion', 6),
                    'extraccion_prom_dia_m3s': prom('penitas', 'extraccion'),
                    'aportado_prom_6h_m3s': prom('penitas', 'aportado', 6)},
        'juan_grijalva': {'nivel_msnm': r2(jg['nivel']), 'aportado_m3s': r2(jg['aportado'], 1),
                          'canal_m3s': r2(jg['canal'], 1), 'tuneles_m3s': r2(jg['tuneles'], 1),
                          'salida_total_m3s': r2(jg['total'], 1), 'salida_prom_6h_m3s': prom('juan_grijalva', 'total', 6)},
        'malpaso': {'nivel_msnm': r2(mp['nivel']), 'namo_msnm': NAMO['Malpaso'],
                    'dif_namo_m': r2(NAMO['Malpaso'] - mp['nivel']) if mp['nivel'] is not None else None,
                    'extraccion_m3s': r2(mp['extraccion'], 1), 'aportado_m3s': r2(apm.get(h), 1)},
        'angostura': {'nivel_msnm': r2((niv.get(h) or {}).get('angostura'))},
        'chicoasen': {'nivel_msnm': r2((niv.get(h) or {}).get('chicoasen')), 'namo_msnm': NAMO['Chicoasén']},
        'novedades': novedades,
        'horario': horario,
        'nota': 'Valores horarios de operación de presas tomados del reporte de CFE; sólo se publican los que usa el modelo.',
    }


def load_state():
    try:
        return json.loads(STATE.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}


def save_state(st):
    OUT.mkdir(parents=True, exist_ok=True)
    STATE.write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding='utf-8')


def dec(s):
    try:
        return str(make_header(decode_header(s or '')))
    except Exception:
        return s or ''


def main():
    user = os.environ.get('GMAIL_BACKUP_USER', '')
    pwd = os.environ.get('GMAIL_BACKUP_APP_PASSWORD', '')
    st = load_state()
    st['revision_local'] = now_local().isoformat(timespec='minutes')
    if not user or not pwd:
        st['error'] = 'sin credenciales Gmail'
        save_state(st)
        return
    mail = None
    try:
        for attempt in range(1, 4):
            try:
                mail = imaplib.IMAP4_SSL('imap.gmail.com', 993, timeout=30)
                mail.login(user, pwd)
                mail.select('INBOX', readonly=True)
                break
            except (OSError, imaplib.IMAP4.abort, imaplib.IMAP4.error) as exc:
                mail = None
                if attempt == 3:
                    raise
                print(f'IMAP intento {attempt}/3 falló: {exc}')
                time.sleep(attempt * 5)
        since = (now_local() - timedelta(days=2)).strftime('%d-%b-%Y')
        typ, data = mail.uid('search', None, f'(SINCE "{since}" FROM "cfe.mx" SUBJECT "FUNCIONAMIENTO DE VASOS")')
        uids = [int(u) for u in (data[0] or b'').split()] if typ == 'OK' else []
        if not uids:
            st['error'] = None
            st['resultado'] = 'sin correos CFE recientes'
            save_state(st)
            print('CFE vasos: sin correos recientes.')
            return
        # El más reciente por fecha del mensaje (sólo encabezados; el adjunto se baja una vez).
        best = None
        for uid in uids[-12:]:
            typ, hd = mail.uid('fetch', str(uid), '(BODY.PEEK[HEADER.FIELDS (DATE MESSAGE-ID)])')
            if typ != 'OK' or not hd or not isinstance(hd[0], tuple):
                continue
            h = email.message_from_bytes(hd[0][1])
            try:
                d = parsedate_to_datetime(h.get('Date'))
            except (TypeError, ValueError):
                continue
            if best is None or d > best[1]:
                best = (uid, d, (h.get('Message-ID') or '').strip())
        if best is None:
            raise RuntimeError('no se pudieron leer las fechas de los correos CFE')
        uid, msg_dt, mid = best
        if mid and mid == st.get('message_id'):
            st['error'] = None
            st['resultado'] = 'sin correo nuevo'
            save_state(st)
            print('CFE vasos: el correo más reciente ya estaba procesado.')
            return
        typ, md = mail.uid('fetch', str(uid), '(BODY.PEEK[])')
        if typ != 'OK' or not md or not isinstance(md[0], tuple):
            raise RuntimeError('no se pudo descargar el correo CFE')
        msg = email.message_from_bytes(md[0][1])
        att = None
        for part in msg.walk():
            fn = dec(part.get_filename())
            if fn and NAME_RE.match(fn.strip()):
                att = (fn.strip(), part.get_payload(decode=True))
                break
        if not att or not att[1]:
            raise RuntimeError('el correo CFE no trae adjunto FINddmmaa.xlsx')
        out = build(att[1], att[0], msg_dt)
        OUT.mkdir(parents=True, exist_ok=True)
        (OUT / 'historico').mkdir(exist_ok=True)
        txt = json.dumps(out, ensure_ascii=False, indent=1)
        prev = None
        try:
            prev = json.loads((OUT / 'latest.json').read_text(encoding='utf-8'))
        except (OSError, ValueError):
            pass
        # No retroceder: sólo se publica si el dato es igual o más reciente que el vigente.
        if not prev or (out['fecha_hora_ultimo_dato'] >= (prev.get('fecha_hora_ultimo_dato') or '')):
            (OUT / 'latest.json').write_text(txt, encoding='utf-8')
        (OUT / 'historico' / f"{out['fecha']}.json").write_text(txt, encoding='utf-8')
        st.update({'message_id': mid, 'mensaje_local': out['mensaje_local'], 'archivo': att[0],
                   'ultimo_dato': out['fecha_hora_ultimo_dato'], 'error': None, 'resultado': 'procesado'})
        save_state(st)
        p = out['penitas']
        print(f"CFE vasos {att[0]}: Peñitas {p['nivel_msnm']} msnm, extracción {p['extraccion_m3s']} m3/s "
              f"(prom 6 h {p['extraccion_prom_6h_m3s']}), dato {out['hora_ultimo_dato']} h.")
    except Exception as exc:  # nunca rompe el flujo de Gmail
        st['error'] = f'{type(exc).__name__}: {exc}'[:300]
        save_state(st)
        print('CFE vasos: error', st['error'])
    finally:
        if mail is not None:
            try:
                mail.logout()
            except Exception:
                pass


if __name__ == '__main__':
    main()
